"""Deterministic ATS-safe DOCX/PDF rendering for validated v2 application drafts."""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
from pathlib import PurePath
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo
import hashlib
import re

from docx import Document
from docx.shared import Inches, Pt
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen.canvas import Canvas

from .artifacts import ARTIFACT_MEDIA_TYPES, ApplicationArtifact, artifact_id_for
from .preparations import ApplicationPreparation, DRAFT_RESULT_SCHEMA_V2


_SAFE_PART = re.compile(r"[^A-Za-z0-9]+")
_FIXED_DOC_TIME = datetime(2000, 1, 1, 0, 0, 0)


def _slug(value: str, *, fallback: str) -> str:
    text = _SAFE_PART.sub("-", str(value or "").strip()).strip("-")
    return (text[:60] or fallback).strip("-")


def _evidence_rows(snapshot: dict[str, Any]) -> list[tuple[str, dict[str, Any]]]:
    facts = snapshot.get("facts") if isinstance(snapshot, dict) else {}
    if not isinstance(facts, dict):
        return []
    rows: list[tuple[str, dict[str, Any]]] = []
    for fact_type, items in facts.items():
        if not isinstance(items, list):
            continue
        for item in items:
            if isinstance(item, dict) and item.get("fact_id") and item.get("value"):
                rows.append((str(fact_type), dict(item)))
    return rows


def _header_values(snapshot: dict[str, Any], header_ids: list[str]) -> tuple[str, list[str]]:
    allowed = set(header_ids)
    rows = [
        (fact_type, row)
        for fact_type, row in _evidence_rows(snapshot)
        if str(row.get("fact_id")) in allowed and fact_type in {"identity", "contact", "location"}
    ]
    identities = [str(row["value"]) for fact_type, row in rows if fact_type == "identity"]
    contacts = [
        str(row["value"])
        for fact_type, row in rows
        if fact_type in {"contact", "location"}
    ]
    name = identities[0] if identities else "Candidate"
    return name, contacts


def _claim_texts(value: Any) -> list[str]:
    if not isinstance(value, list):
        return []
    result: list[str] = []
    for item in value:
        if isinstance(item, dict):
            text = str(item.get("text") or "").strip()
            if text:
                result.append(text)
    return result


def _normalise_docx_bytes(raw: bytes) -> bytes:
    source = BytesIO(raw)
    output = BytesIO()
    with ZipFile(source, "r") as src, ZipFile(output, "w", compression=ZIP_DEFLATED) as dst:
        for name in sorted(src.namelist()):
            data = src.read(name)
            info = ZipInfo(filename=name, date_time=(2000, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            dst.writestr(info, data)
    return output.getvalue()


def _docx_base() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10.5)
    props = doc.core_properties
    props.author = "SmartPloy"
    props.last_modified_by = "SmartPloy"
    props.created = _FIXED_DOC_TIME
    props.modified = _FIXED_DOC_TIME
    props.title = ""
    return doc


def _docx_bytes(doc: Document) -> bytes:
    buffer = BytesIO()
    doc.save(buffer)
    return _normalise_docx_bytes(buffer.getvalue())


def _add_docx_heading(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_paragraph()
    paragraph.style = doc.styles["Normal"]
    run = paragraph.add_run(text)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(14 if level == 1 else 11)


def _build_cv_docx(
    *,
    cv: dict[str, Any],
    evidence: dict[str, Any],
) -> bytes:
    doc = _docx_base()
    name, contacts = _header_values(evidence, list(cv.get("header_evidence_fact_ids") or ()))
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(name)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(18)
    if contacts:
        doc.add_paragraph(" | ".join(contacts))

    headline = str((cv.get("headline") or {}).get("text") or "").strip()
    if headline:
        p = doc.add_paragraph()
        run = p.add_run(headline)
        run.bold = True

    summary = str((cv.get("professional_summary") or {}).get("text") or "").strip()
    if summary:
        _add_docx_heading(doc, "Professional Summary", 2)
        doc.add_paragraph(summary)

    sections = (
        ("skills", "Skills"),
        ("experience", "Experience"),
        ("education", "Education"),
        ("certifications", "Certifications"),
        ("projects", "Projects"),
        ("languages", "Languages"),
    )
    for key, title in sections:
        values = _claim_texts(cv.get(key))
        if not values:
            continue
        _add_docx_heading(doc, title, 2)
        if key == "skills":
            doc.add_paragraph(", ".join(values))
        else:
            for value in values:
                doc.add_paragraph(value, style="List Bullet")

    return _docx_bytes(doc)


def _build_letter_docx(
    *,
    letter: dict[str, Any],
    evidence: dict[str, Any],
    job: dict[str, Any],
) -> bytes:
    doc = _docx_base()
    header_ids = list(letter.get("header_evidence_fact_ids") or ())
    name, contacts = _header_values(evidence, header_ids)
    p = doc.add_paragraph()
    run = p.add_run(name)
    run.bold = True
    run.font.size = Pt(14)
    if contacts:
        doc.add_paragraph(" | ".join(contacts))

    role = str(job.get("role_title") or "").strip()
    company = str(job.get("company") or "").strip()
    if role or company:
        doc.add_paragraph("Application: " + " — ".join(item for item in (role, company) if item))

    salutation = str(letter.get("salutation") or "Dear Hiring Team,").strip()
    doc.add_paragraph(salutation)
    for paragraph in _claim_texts(letter.get("paragraphs")):
        doc.add_paragraph(paragraph)

    closing = str(letter.get("closing") or "Kind regards,").strip()
    doc.add_paragraph(closing)
    doc.add_paragraph(name)
    return _docx_bytes(doc)


def _wrap_line(text: str, *, width: float, font_name: str, font_size: float) -> list[str]:
    words = str(text or "").split()
    if not words:
        return [""]
    lines: list[str] = []
    current = words[0]
    for word in words[1:]:
        candidate = current + " " + word
        if stringWidth(candidate, font_name, font_size) <= width:
            current = candidate
        else:
            lines.append(current)
            current = word
    lines.append(current)
    return lines


class _PdfWriter:
    def __init__(self) -> None:
        self.buffer = BytesIO()
        self.canvas = Canvas(self.buffer, pagesize=A4, pageCompression=1, invariant=1)
        self.width, self.height = A4
        self.left = 48
        self.right = 48
        self.top = 50
        self.bottom = 50
        self.y = self.height - self.top

    def _page_if_needed(self, line_height: float) -> None:
        if self.y - line_height < self.bottom:
            self.canvas.showPage()
            self.y = self.height - self.top

    def text(
        self,
        value: str,
        *,
        size: float = 10.5,
        bold: bool = False,
        gap_after: float = 5,
        bullet: bool = False,
    ) -> None:
        font = "Helvetica-Bold" if bold else "Helvetica"
        prefix = "• " if bullet else ""
        line_height = size * 1.35
        available = self.width - self.left - self.right
        lines = _wrap_line(prefix + str(value or ""), width=available, font_name=font, font_size=size)
        for line in lines:
            self._page_if_needed(line_height)
            self.canvas.setFont(font, size)
            self.canvas.drawString(self.left, self.y, line)
            self.y -= line_height
        self.y -= gap_after

    def finish(self) -> bytes:
        self.canvas.save()
        return self.buffer.getvalue()


def _build_cv_pdf(*, cv: dict[str, Any], evidence: dict[str, Any]) -> bytes:
    pdf = _PdfWriter()
    name, contacts = _header_values(evidence, list(cv.get("header_evidence_fact_ids") or ()))
    pdf.text(name, size=17, bold=True, gap_after=2)
    if contacts:
        pdf.text(" | ".join(contacts), size=9.5, gap_after=8)

    headline = str((cv.get("headline") or {}).get("text") or "").strip()
    if headline:
        pdf.text(headline, size=11.5, bold=True, gap_after=8)

    summary = str((cv.get("professional_summary") or {}).get("text") or "").strip()
    if summary:
        pdf.text("Professional Summary", size=11.5, bold=True, gap_after=3)
        pdf.text(summary, gap_after=8)

    sections = (
        ("skills", "Skills"),
        ("experience", "Experience"),
        ("education", "Education"),
        ("certifications", "Certifications"),
        ("projects", "Projects"),
        ("languages", "Languages"),
    )
    for key, title in sections:
        values = _claim_texts(cv.get(key))
        if not values:
            continue
        pdf.text(title, size=11.5, bold=True, gap_after=3)
        if key == "skills":
            pdf.text(", ".join(values), gap_after=8)
        else:
            for value in values:
                pdf.text(value, bullet=True, gap_after=2)
            pdf.y -= 5
    return pdf.finish()


def _build_letter_pdf(
    *,
    letter: dict[str, Any],
    evidence: dict[str, Any],
    job: dict[str, Any],
) -> bytes:
    pdf = _PdfWriter()
    name, contacts = _header_values(
        evidence,
        list(letter.get("header_evidence_fact_ids") or ()),
    )
    pdf.text(name, size=14, bold=True, gap_after=2)
    if contacts:
        pdf.text(" | ".join(contacts), size=9.5, gap_after=10)

    role = str(job.get("role_title") or "").strip()
    company = str(job.get("company") or "").strip()
    if role or company:
        pdf.text(
            "Application: " + " — ".join(item for item in (role, company) if item),
            bold=True,
            gap_after=10,
        )

    pdf.text(str(letter.get("salutation") or "Dear Hiring Team,"), gap_after=8)
    for paragraph in _claim_texts(letter.get("paragraphs")):
        pdf.text(paragraph, gap_after=8)
    pdf.text(str(letter.get("closing") or "Kind regards,"), gap_after=3)
    pdf.text(name)
    return pdf.finish()


def _artifact(
    preparation: ApplicationPreparation,
    *,
    artifact_type: str,
    filename: str,
    content: bytes,
) -> ApplicationArtifact:
    media_type = ARTIFACT_MEDIA_TYPES[artifact_type][0]
    digest = hashlib.sha256(content).hexdigest()
    return ApplicationArtifact(
        artifact_id=artifact_id_for(
            package_id=preparation.package_id,
            artifact_type=artifact_type,
            content_sha256=digest,
        ),
        package_id=preparation.package_id,
        canonical_key=preparation.canonical_key,
        artifact_type=artifact_type,
        filename=filename,
        media_type=media_type,
        content_sha256=digest,
        size_bytes=len(content),
        content=content,
        created_at=preparation.completed_at or preparation.updated_at,
        metadata={
            "draft_schema": str(preparation.result.get("schema") or ""),
            "preparation_contract": preparation.preparation_contract,
            "evidence_sha256": preparation.evidence_sha256,
            "job_snapshot_sha256": preparation.job_snapshot_sha256,
        },
    ).normalized()


def render_application_artifacts(
    preparation: ApplicationPreparation,
) -> list[ApplicationArtifact]:
    item = preparation.normalized()
    if item.status != "completed":
        raise ValueError("application documents require a completed preparation")
    if item.result.get("schema") != DRAFT_RESULT_SCHEMA_V2:
        return []

    cv = item.result.get("tailored_cv")
    letter = item.result.get("cover_letter")
    if not isinstance(cv, dict) or not isinstance(letter, dict):
        raise ValueError("validated v2 result is missing document structures")

    company = _slug(str(item.job_snapshot.get("company") or ""), fallback="Employer")
    role = _slug(str(item.job_snapshot.get("role_title") or ""), fallback="Role")
    base = f"{company}_{role}"

    cv_docx = _build_cv_docx(cv=cv, evidence=item.evidence_snapshot)
    cv_pdf = _build_cv_pdf(cv=cv, evidence=item.evidence_snapshot)
    letter_docx = _build_letter_docx(
        letter=letter,
        evidence=item.evidence_snapshot,
        job=item.job_snapshot,
    )
    letter_pdf = _build_letter_pdf(
        letter=letter,
        evidence=item.evidence_snapshot,
        job=item.job_snapshot,
    )

    return [
        _artifact(item, artifact_type="cv_docx", filename=f"{base}_CV.docx", content=cv_docx),
        _artifact(item, artifact_type="cv_pdf", filename=f"{base}_CV.pdf", content=cv_pdf),
        _artifact(
            item,
            artifact_type="cover_letter_docx",
            filename=f"{base}_Cover-Letter.docx",
            content=letter_docx,
        ),
        _artifact(
            item,
            artifact_type="cover_letter_pdf",
            filename=f"{base}_Cover-Letter.pdf",
            content=letter_pdf,
        ),
    ]
