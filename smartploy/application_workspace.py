"""Locally assembled, review-first application documents (no AI or employer actions)."""
from __future__ import annotations

from dataclasses import dataclass, replace
from hashlib import sha256
from io import BytesIO
from html import escape
import re
from typing import Literal

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

from .domain import JobCandidate, normalize_identity_text, utc_now_iso
from .profile import CandidateProfile
from .scoring import score_candidate
from .complete_letter import compose_completed_letter

MAX_CV_CHARACTERS = 30000
MAX_LETTER_CHARACTERS = 10000
MAX_DOCUMENT_BYTES = 4 * 1024 * 1024


class WorkspaceError(ValueError):
    pass


class WorkspaceRevisionConflict(WorkspaceError):
    pass


def _clean_text(value: str, *, limit: int, label: str) -> str:
    if not isinstance(value, str) or not value.strip() or len(value) > limit:
        raise WorkspaceError(f"{label} must contain 1 to {limit} characters")
    if any(ord(char) < 32 and char not in "\n\t\r" for char in value):
        raise WorkspaceError(f"{label} contains unsupported control characters")
    return value.replace("\r\n", "\n").replace("\r", "\n").strip()


def _safe_excerpt(value: str, maximum: int) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()[:maximum]


def _job_label(value: str, fallback: str) -> str:
    cleaned = _safe_excerpt(value, 150)
    return cleaned if cleaned else fallback


def _candidate_document(profile: CandidateProfile) -> str:
    """Render approved facts verbatim; never infer qualifications or employers."""
    approved = profile.approved_evidence().get("facts", {})
    lines: list[str] = []
    for type_name, heading in (
        ("identity", "NAME"), ("contact", "CONTACT"),
        ("experience", "EXPERIENCE"), ("education", "EDUCATION"),
        ("skill", "SKILLS"), ("certification", "CERTIFICATIONS"),
        ("project", "PROJECTS"), ("language", "LANGUAGES"),
        ("domain", "SPECIALISATIONS"), ("work_authorization", "WORK AUTHORISATION"),
        ("location", "LOCATION"),
    ):
        items = [str(row.get("value") or "").strip() for row in approved.get(type_name, [])]
        items = [value for value in items if value]
        if items:
            lines.extend([heading, *items, ""])
    if not lines:
        raise WorkspaceError("Upload a CV or approve profile facts before preparing an application")
    return "\n".join(lines).strip()


# The existing applicant evidence scorer already checks declared job
# requirements and known skills against exact APPROVED profile fact IDs.
# Reuse it instead of writing a second (less trustworthy) match engine.
# Profile facts are free text, so show them verbatim rather than inventing
# achievements, durations or qualification levels around a match.
_LETTER_FACT_LABELS: dict[str, tuple[str, str]] = {
    "skill": ("Skill", "Osaaminen"),
    "experience": ("Experience", "Kokemus"),
    "education": ("Education", "Koulutus"),
    "project": ("Project", "Projekti"),
    "certification": ("Certification", "Sertifikaatti"),
    "domain": ("Specialisation", "Erityisosaaminen"),
    "language": ("Language", "Kielitaito"),
}
_LETTER_FACT_PRIORITY = {
    "skill": 0, "experience": 1, "project": 2,
    "certification": 3, "education": 4, "domain": 5, "language": 6,
}
_FINNISH_JOB_LANGUAGES = frozenset({
    "finnish", "suomi", "suomen kieli", "suomenkielinen",
    "fi", "fi fi", "finska",
})


def _relevant_approved_letter_facts(
    candidate: JobCandidate, profile: CandidateProfile,
) -> tuple[tuple[str, str], ...]:
    """Choose up to three ORIGINAL approved facts relevant to THIS posting.

    Score IDs are not professional credentials. No language, degree or claim
    can be inferred from the job description; only exact approved facts are
    considered for the applicant's draft.
    """
    approved = profile.approved_facts
    matched_ids = set(score_candidate(candidate, profile).matched_evidence_ids)
    candidates = [
        (fact.fact_type, fact.value, fact.fact_id)
        for fact in approved
        if fact.fact_type in _LETTER_FACT_LABELS
        and fact.fact_id in matched_ids
        # Never truncate a factual statement: removing the trailing negation,
        # limit or date could invert its meaning.
        and 1 <= len(fact.value) <= 260
    ]
    candidates.sort(
        key=lambda item: (_LETTER_FACT_PRIORITY[item[0]], item[2])
    )
    return tuple((kind, value) for kind, value, _ in candidates[:3])


def cover_letter_proposal(
    candidate: JobCandidate, profile: CandidateProfile,
    *, language: str | None = None,
) -> str:
    """Produce a reviewable, job-aligned draft without remote AI or invented claims.

    Only an EXPLICIT published Finnish-language marker or applicant-requested
    language selects Finnish; unsupported or mixed markers default to English.
    The name is used only when its original identity fact was approved.
    """
    role = _job_label(candidate.job.role_title, "the advertised role")
    company = _job_label(candidate.job.company, "your organisation")
    requested = normalize_identity_text(
        candidate.job.language if language is None else language
    )
    finnish = requested in _FINNISH_JOB_LANGUAGES
    evidence = profile.approved_evidence().get("facts", {})
    name = next(
        (str(f.get("value")) for f in evidence.get("identity", []) if f.get("value")),
        "[Oma nimi]" if finnish else "[Your name]",
    )
    matched = _relevant_approved_letter_facts(candidate, profile)
    if finnish:
        facts = (
            "Haluan nostaa esiin seuraavat tehtävään liittyvät, itse "
            "vahvistamani tiedot:\n"
            + "\n".join(
                f"• {_LETTER_FACT_LABELS[kind][1]}: {value}"
                for kind, value in matched
            )
            if matched else
            "[Lisää tähän tehtävään liittyvä, paikkansapitävä osaamis- "
            "tai työkokemusesimerkki. Älä ilmoita pätevyyttä ilman näyttöä.]"
        )
        return (
            "Hei rekrytointitiimi,\n\n"
            f"Haen avoinna olevaa tehtäväänne ({role}) organisaatiossa {company}.\n\n"
            f"{facts}\n\n"
            "[Kerro omin sanoin, miksi juuri tämä tehtävä ja työnantaja "
            "kiinnostavat sinua. Lisää konkreettinen esimerkki ja tarkista "
            "kaikki väitteet ennen hakemuksen lähettämistä.]\n\n"
            "Kiitos hakemukseni käsittelystä. Kerron mielelläni lisää "
            "osaamisestani haastattelussa.\n\n"
            f"Ystävällisin terveisin,\n{name}"
        )
    facts = (
        "I would like to highlight these relevant parts of my background:\n"
        + "\n".join(
            f"• {_LETTER_FACT_LABELS[kind][0]}: {value}"
            for kind, value in matched
        )
        if matched else
        "[Add a truthful, specific skill or project example relevant to "
        "this role. Do not claim qualifications you cannot substantiate.]"
    )
    return (
        "Dear Hiring Team,\n\n"
        f"I am applying for the {role} position at {company}.\n\n"
        f"{facts}\n\n"
        "[Explain in your own words why this employer and role interest you. "
        "Add a specific example and check all factual statements before "
        "submitting the application.]\n\n"
        "Thank you for considering my application. I would welcome "
        "the opportunity to discuss the role.\n\n"
        f"Kind regards,\n{name}"
    )


@dataclass(frozen=True)
class ApplicationWorkspace:
    workspace_id: str
    profile_id: str
    canonical_key: str
    role_title: str
    company: str
    application_url: str
    cv_text: str
    cover_letter_text: str
    source_cv_sha256: str
    approved_evidence_ids: tuple[str, ...]
    job_fingerprint: str
    revision: int = 1
    status: str = "draft"
    self_reported_applied_at: str = ""
    self_reported_reference: str = ""
    created_at: str = ""
    updated_at: str = ""

    def to_dict(self) -> dict:
        return {
            "workspace_id": self.workspace_id,
            "canonical_key": self.canonical_key,
            "role_title": self.role_title,
            "company": self.company,
            "application_url": self.application_url,
            "cv_text": self.cv_text,
            "cover_letter_text": self.cover_letter_text,
            "source_cv_sha256": self.source_cv_sha256,
            "approved_evidence_ids": list(self.approved_evidence_ids),
            "job_fingerprint": self.job_fingerprint,
            "revision": self.revision,
            "status": self.status,
            "self_reported_applied_at": self.self_reported_applied_at,
            "self_reported_reference": self.self_reported_reference,
            "created_at": self.created_at,
            "updated_at": self.updated_at,
            "authority": {
                "may_edit": self.status == "draft",
                "may_download": True,
                "may_submit_to_employer": False,
                # Official provider application ID is stored in a separate
                # owner-scoped receipt ledger, never in self-report fields.
                "has_external_submission_proof": self.status == "provider_acknowledged",
            },
        }


def new_workspace(
    candidate: JobCandidate, profile: CandidateProfile,
    source_cv_text: str = "", source_cv_sha256: str = "",
) -> ApplicationWorkspace:
    if not candidate.source_current:
        raise WorkspaceError("This job is no longer current. Check the employer posting.")
    source = source_cv_text.strip()
    if source and len(source) > MAX_CV_CHARACTERS:
        raise WorkspaceError(
            "Your uploaded CV contains more than 30,000 characters. Please upload a shorter CV."
        )
    cv_text = source if source else _candidate_document(profile)
    cv_text = _clean_text(cv_text, limit=MAX_CV_CHARACTERS, label="CV")
    # SP85: a new application opens with ACTUAL bilingual letter prose, not
    # the original coaching template. User-approved job-matched evidence alone
    # supplies applicant claims; missing facts never become fabricated text.
    # Existing persisted workspaces are never changed by this default.
    initial_language = (
        "fi" if normalize_identity_text(candidate.job.language) in _FINNISH_JOB_LANGUAGES
        else "en"
    )
    letter = _clean_text(
        compose_completed_letter(candidate, profile, language=initial_language).text,
        limit=MAX_LETTER_CHARACTERS, label="Cover letter",
    )
    pid = profile.profile_id
    identity = sha256((pid + "\0" + candidate.canonical_key).encode("utf-8")).hexdigest()[:30]
    approved = profile.approved_evidence().get("facts", {})
    ids = tuple(sorted({
        str(f["fact_id"]) for rows in approved.values() for f in rows if f.get("fact_id")
    }))
    now = utc_now_iso()
    return ApplicationWorkspace(
        workspace_id="kit-" + identity,
        profile_id=pid, canonical_key=candidate.canonical_key,
        role_title=_job_label(candidate.job.role_title, "Job application"),
        company=_job_label(candidate.job.company, "Employer"),
        application_url=str(candidate.job.application_url or ""),
        cv_text=cv_text, cover_letter_text=letter,
        source_cv_sha256=source_cv_sha256 if source else "",
        approved_evidence_ids=ids,
        job_fingerprint=candidate.description_fingerprint,
        created_at=now, updated_at=now,
    )


def update_documents(workspace: ApplicationWorkspace, *, cv_text: str,
                     cover_letter_text: str) -> ApplicationWorkspace:
    if workspace.status != "draft":
        raise WorkspaceError("An applied workspace cannot be edited. Create a new application separately.")
    return replace(
        workspace,
        cv_text=_clean_text(cv_text, limit=MAX_CV_CHARACTERS, label="CV"),
        cover_letter_text=_clean_text(cover_letter_text, limit=MAX_LETTER_CHARACTERS, label="Cover letter"),
        revision=workspace.revision + 1, updated_at=utc_now_iso(),
    )


def mark_applied(workspace: ApplicationWorkspace, *, reference: str = "") -> ApplicationWorkspace:
    if workspace.status == "self_reported_applied":
        return workspace
    if workspace.status != "draft":
        raise WorkspaceError("A provider-acknowledged application cannot be downgraded to self-reported")
    ref = str(reference or "").strip()
    if len(ref) > 250 or any(ord(char) < 32 for char in ref):
        raise WorkspaceError("Application reference must be 250 printable characters or fewer")
    now = utc_now_iso()
    return replace(
        workspace, status="self_reported_applied", self_reported_applied_at=now,
        self_reported_reference=ref, revision=workspace.revision + 1, updated_at=now,
    )


def _slug(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", value.casefold()).strip("-")[:48] or "application"


def download_filename(workspace: ApplicationWorkspace, kind: str, file_format: str) -> str:
    if kind not in {"cv", "letter"} or file_format not in {"pdf", "docx"}:
        raise WorkspaceError("Unsupported document kind or format")
    return f"{_slug(workspace.company)}-{_slug(workspace.role_title)}-{kind}.{file_format}"


def render_document(
    workspace: ApplicationWorkspace,
    *, kind: Literal["cv", "letter"], file_format: Literal["pdf", "docx"],
) -> tuple[bytes, str, str]:
    """Generate ATS-friendly user-editable content; no network or employer submission."""
    filename = download_filename(workspace, kind, file_format)
    raw = workspace.cv_text if kind == "cv" else workspace.cover_letter_text
    text = _clean_text(
        raw, label="CV" if kind == "cv" else "Cover letter",
        limit=MAX_CV_CHARACTERS if kind == "cv" else MAX_LETTER_CHARACTERS,
    )
    heading = "Curriculum Vitae" if kind == "cv" else f"Application — {workspace.role_title}"
    output = BytesIO()
    if file_format == "docx":
        doc = Document()
        section = doc.sections[0]
        section.top_margin = Inches(.8)
        section.bottom_margin = Inches(.75)
        section.left_margin = Inches(.9)
        section.right_margin = Inches(.9)
        normal = doc.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(10.5)
        normal.font.color.rgb = RGBColor(31, 49, 75)
        normal.paragraph_format.space_after = Pt(5)
        doc.add_heading(heading, level=1)
        for line in text.split("\n"):
            doc.add_paragraph(line)
        doc.save(output)
        mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    else:
        style = ParagraphStyle(
            "Body", fontName="Helvetica", fontSize=10.2, leading=14.8,
            textColor=colors.HexColor("#20334A"), alignment=TA_LEFT,
            spaceAfter=7, splitLongWords=1,
        )
        title_style = ParagraphStyle(
            "Heading", parent=style, fontName="Helvetica-Bold", fontSize=15,
            leading=20, spaceAfter=15, textColor=colors.HexColor("#0D3262"),
        )
        story = [Paragraph(escape(heading), title_style)]
        for line in text.split("\n"):
            if line.strip():
                story.append(Paragraph(escape(line), style))
            else:
                story.append(Spacer(1, 5))
        pdf = SimpleDocTemplate(
            output, pagesize=A4, topMargin=21*mm, bottomMargin=21*mm,
            leftMargin=23*mm, rightMargin=23*mm,
            title=heading, author="SmartPloy user-reviewed application", invariant=1,
        )
        pdf.build(story)
        mime = "application/pdf"
    content = output.getvalue()
    if not content or len(content) > MAX_DOCUMENT_BYTES:
        raise WorkspaceError("Generated document exceeds safe download limits")
    return content, mime, filename
