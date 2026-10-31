"""SP79: an editable, versioned CV that really feeds existing matching/workspaces.

No external employer request, automatic approval or model calls in this module.
"""
from __future__ import annotations

from base64 import b64encode
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
import os
from uuid import uuid4

from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfReader
import pytest

from smartploy.api import create_app
from smartploy.cv_editor import render_editor_cv
from smartploy.cv_storage import (
    CVRevisionConflict, MemoryCVRepository, PostgresCVRepository,
)
from smartploy.service import SmartployService
from smartploy.storage import PostgresJobRepository


DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def docx(data: str) -> bytes:
    document = Document()
    document.add_paragraph(data)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def test_master_cv_can_be_created_without_upload_then_saved_and_exported():
    repo = MemoryCVRepository()
    text = "Alex Applicant\n\nSKILLS\nPython and data analytics."
    saved = repo.save_editor_text("person-a", expected_sha256="", text=text)
    assert saved.filename == "smartploy-master-cv.docx"
    assert saved.summary()["created_by_editor"] is True
    assert saved.text == text
    assert repo.history("person-a") == []
    content = Document(BytesIO(saved.content))
    assert "Alex Applicant" in "\n".join(p.text for p in content.paragraphs)
    assert repo.get("person-b") is None
    pdf = render_editor_cv(saved.text, file_format="pdf")
    assert pdf.startswith(b"%PDF-")
    assert "Alex Applicant" in "".join(page.extract_text() or "" for page in PdfReader(BytesIO(pdf)).pages)


def test_editor_backups_original_and_blocks_stale_write_or_other_owner():
    repo = MemoryCVRepository()
    original = docx("Alex Applicant. Working with SQL.")
    uploaded = repo.put("person-a", "original.docx", DOCX_MIME,
                        original, "Alex Applicant. Working with SQL.")
    saved = repo.save_editor_text(
        "person-a", expected_sha256=uploaded.digest,
        text="Alex Applicant\n\nSKILLS\nPython, SQL and Power BI.",
    )
    assert saved.digest != uploaded.digest
    assert repo.get_archived("person-a", uploaded.digest).content == original
    assert repo.get_archived("person-b", uploaded.digest) is None
    assert repo.history("person-a")[0]["sha256"] == uploaded.digest
    with pytest.raises(CVRevisionConflict):
        repo.save_editor_text(
            "person-a", expected_sha256=uploaded.digest,
            text="A stale tab must never overwrite the current application CV.",
        )
    newer = repo.save_editor_text(
        "person-a", expected_sha256=saved.digest,
        text="Alex Applicant\n\nSKILLS\nPython, SQL, Power BI and testing.",
    )
    assert newer.text.endswith("testing.")
    assert len(repo.history("person-a")) == 2
    assert repo.delete("person-a")
    assert repo.get("person-a") is None
    assert repo.get_archived("person-a", uploaded.digest) is None


@pytest.mark.parametrize("value", [
    "", "short", "Valid CV text with NUL \x00 and hidden payload",
    "A" * 30_001,
])
def test_invalid_master_cv_is_rejected_without_partial_save(value):
    repo = MemoryCVRepository()
    with pytest.raises(ValueError):
        repo.save_editor_text("person-a", expected_sha256="", text=value)
    assert repo.get("person-a") is None


def test_editor_api_is_actual_document_workflow_not_a_drafting_guide():
    api = TestClient(create_app(SmartployService()))
    missing = api.get("/api/v1/profile/cv/editor")
    assert missing.status_code == 200
    assert missing.json()["present"] is False
    assert missing.headers["cache-control"] == "private, no-store"
    text = "Synthetic Applicant\n\nEXPERIENCE\nTest-driven quality assurance and Python."
    saving = api.put("/api/v1/profile/cv/editor", json={
        "expected_sha256": "", "text": text, "confirmed_replace": True,
    })
    assert saving.status_code == 200, saving.text
    state = saving.json()
    assert state["cv"]["created_by_editor"] is True
    assert state["employer_actions_performed"] == 0
    assert state["approved_profile_facts_added"] == 0
    current = api.get("/api/v1/profile/cv/editor")
    assert current.json()["text"] == text
    assert len(current.json()["sha256"]) == 64
    assert current.headers["cache-control"] == "private, no-store"
    assert api.get("/api/v1/profile/cv").json()["text_preview"] == text
    assert api.get("/api/v1/profile/status").json()["approved_fact_count"] == 0
    for extension, prefix in (("pdf", b"%PDF-"), ("docx", b"PK")):
        result = api.get("/api/v1/profile/cv/editor/download/" + extension)
        assert result.status_code == 200, result.text
        assert result.content.startswith(prefix)
        assert result.headers["cache-control"] == "private, no-store"
        assert "attachment" in result.headers["content-disposition"]
    assert api.get("/api/v1/profile/cv/editor/download/html").status_code == 422
    assert api.put("/api/v1/profile/cv/editor", json={
        "expected_sha256": "", "text": text + " More.",
        "confirmed_replace": True,
    }).status_code == 409
    assert api.put("/api/v1/profile/cv/editor", json={
        "expected_sha256": current.json()["sha256"],
        "text": text + " More.", "confirmed_replace": False,
    }).status_code == 422
    assert api.put("/api/v1/profile/cv/editor", json={
        "expected_sha256": current.json()["sha256"],
        "text": text, "confirmed_replace": True, "submit_to_employer": True,
    }).status_code == 422
    assert api.delete("/api/v1/profile/cv").status_code == 200
    assert api.get("/api/v1/profile/cv/history/" + current.json()["sha256"]).status_code == 404


def test_editor_exposes_recoverable_exact_original_after_upload():
    api = TestClient(create_app(SmartployService()))
    original = docx("Alex Applicant has a confirmed real SQL project.")
    sent = api.post("/api/v1/profile/cv", json={
        "filename": "My original.docx",
        "content_base64": b64encode(original).decode("ascii"),
    })
    assert sent.status_code == 200, sent.text
    before = sent.json()
    editor = api.get("/api/v1/profile/cv/editor").json()
    assert editor["sha256"] == before["sha256"]
    assert editor["text"] == "Alex Applicant has a confirmed real SQL project."
    revision = api.put("/api/v1/profile/cv/editor", json={
        "expected_sha256": before["sha256"], "confirmed_replace": True,
        "text": "Alex Applicant\n\nSKILLS\nSQL and Python, reviewed and updated.",
    })
    assert revision.status_code == 200, revision.text
    assert revision.json()["history"][0]["sha256"] == before["sha256"]
    original_download = api.get("/api/v1/profile/cv/history/" + before["sha256"])
    assert original_download.status_code == 200
    assert original_download.content == original
    assert original_download.headers["x-content-type-options"] == "nosniff"
    assert api.get("/api/v1/profile/cv/history/" + ("0" * 64)).status_code == 404
    assert api.delete("/api/v1/profile/cv").json() == {"deleted": True}
    assert api.get("/api/v1/profile/cv/history/" + before["sha256"]).status_code == 404


@pytest.mark.skipif(not os.getenv("SMARTPLOY_TEST_DATABASE_URL"), reason="PostgreSQL integration database required")
def test_postgres_editor_atomic_revision_and_archive_is_owner_scoped():
    dsn = os.environ["SMARTPLOY_TEST_DATABASE_URL"]
    PostgresJobRepository(dsn).migrate()
    repo = PostgresCVRepository(dsn)
    a, b = "sp79-a-" + uuid4().hex, "sp79-b-" + uuid4().hex
    original = docx("Synthetic candidate and original experience in Python.")
    try:
        a0 = repo.put(a, "old.docx", DOCX_MIME, original,
                      "Synthetic candidate and original experience in Python.")
        assert repo.get_archived(b, a0.digest) is None
        with ThreadPoolExecutor(max_workers=2) as pool:
            results = list(pool.map(
                lambda suffix: (
                    repo.save_editor_text(a, expected_sha256=a0.digest,
                        text="Synthetic edited CV and skills in SQL. " + suffix)
                ), ("version A", "version B")
            ))
        # Unreachable: one competing revision MUST conflict.
        raise AssertionError("Both concurrent updates unexpectedly succeeded: " + str(results))
    except CVRevisionConflict:
        current = repo.get(a)
        assert current and current.digest != a0.digest
        assert repo.get_archived(a, a0.digest).content == original
        assert repo.get_archived(b, a0.digest) is None
        assert repo.history(b) == []
        assert repo.history(a)[0]["sha256"] == a0.digest
    finally:
        repo.delete(a)
        repo.delete(b)
